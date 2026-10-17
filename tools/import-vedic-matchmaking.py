"""Import workbook values literally; preserve directional axes and source-review flags."""
from pathlib import Path
import sys,json,hashlib,collections,shutil
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
src=Path(sys.argv[1]);dest=ROOT/'data/vedic-matchmaking-2026-10-17';dest.mkdir(exist_ok=True)
SIGNS=['मेष','वृषभ','मिथुन','कर्क','सिंह','कन्या','तुला','वृश्चिक','धनु','मकर','कुंभ','मीन']
STARS=['अश्वि','भर','कृत्ति','रोहि','मृग','आर्द्रा','पुन','पुष्य','आश्रे','मघा','पूर्वा','उत्त','हस्त','चित्रा','स्वा','विशा','अनु','ज्ये','मूळ','पूषा','उषा','श्रव','धनि','शत','पूभा','उभा','रेव']
def segment(sign,nak,count):
 si=SIGNS.index(str(sign).strip());ni=STARS.index(str(nak).strip())+1
 padas=[p for p in range(1,5) if ((ni-1)*48000+(p-1)*12000)//108000==si]
 if len(padas)!=int(count):raise ValueError(f'Invalid printed pada count: {sign} {nak} {count}; actual {padas}')
 return {'sign':si,'nakshatra':ni,'padas':padas,'signName':sign,'starLabel':nak,'printedPadaCount':int(count)}
w=openpyxl.load_workbook(src,data_only=False);cells=[];keys=set();notes={}
for sheet in w:
 if sheet.sheet_state!='visible':raise ValueError('Unexpected hidden sheet')
 notes[sheet.title]=[sheet.cell(r,1).value for r in range(1,5)]
 for row in range(7,sheet.max_row+1):
  values=[sheet.cell(row,c).value for c in range(1,11)]
  if all(v is None for v in values):continue
  if any(sheet.cell(row,c).data_type=='f' for c in range(1,11)):raise ValueError('Formula found; expected literal source values')
  bs,bn,bp,gs,gn,gp,score,defects,mark,review=values
  bride,groom=segment(bs,bn,bp),segment(gs,gn,gp)
  key=f"{bride['sign']}:{bride['nakshatra']}|{groom['sign']}:{groom['nakshatra']}"
  if key in keys:raise ValueError(f'Duplicate directional cell {key}')
  keys.add(key);marker=str(mark or '').strip();codes=str(defects or '').strip()
  if codes and (not codes.isascii() or not codes.isdigit() or any(d not in '12345678' for d in codes)):raise ValueError(f'Invalid defect codes at {sheet.title}:{row}')
  if marker not in ['', '★','X']:raise ValueError('Unknown source marker')
  if marker and score is not None:raise ValueError('A marked source cell must retain its blank numeric score')
  if not marker and (isinstance(score,bool) or not isinstance(score,(float,int)) or not 0<=score<=36 or score*2%1):raise ValueError('Invalid numeric score')
  cells.append({'key':key,'bride':bride,'groom':groom,'score':score,'defects':[int(d) for d in codes],'defectText':codes,'marker':marker,'review':str(review or ''),'sheet':sheet.title,'row':row})
if len(cells)!=1296:raise ValueError(f'Expected 36 × 36 entries; got {len(cells)}')
segments={f"{c['bride']['sign']}:{c['bride']['nakshatra']}" for c in cells}
if len(segments)!=36:raise ValueError('Expected 36 sign/star segments')
obj={'metadata':{'filename':src.name,'sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'axes':{'rows':'bride','columns':'groom'},'entryCount':len(cells),'segmentCount':36,'notes':notes,'numericScoreCount':sum(c['score'] is not None for c in cells),'markerCounts':dict(collections.Counter(c['marker'] for c in cells)),'policy':'Workbook values imported literally. Blank marked scores remain null. No withdrawn chat rules are reused.'},'cells':cells}
if len(sys.argv)>2:
 nsrc=Path(sys.argv[2]);nw=openpyxl.load_workbook(nsrc,data_only=False);ns=nw['नाडीपादवेध'];groups=[];seen=set()
 aliases={'ज्येष्ठा':'ज्ये','आश्ले':'आश्रे','स्वाती':'स्वा'}
 for r in range(3,15):
  group=[]
  for col in range(1,10):
   cell=ns.cell(r,col)
   if cell.data_type=='f':raise ValueError('Nadi table must contain literal entries')
   label,pada=str(cell.value).rsplit(' ',1);nak=STARS.index(aliases.get(label,label))+1;pada=int(pada)
   key=f'{nak}:{pada}'
   if not 1<=pada<=4 or key in seen:raise ValueError('Invalid or duplicate Nadi pada')
   seen.add(key);group.append({'nakshatra':nak,'pada':pada,'label':cell.value})
  groups.append({'row':r,'entries':group})
 if len(seen)!=108:raise ValueError('Nadi table must cover all 108 padas')
 obj['nadi']={'filename':nsrc.name,'sha256':hashlib.sha256(nsrc.read_bytes()).hexdigest(),'sheet':ns.title,'groups':groups,'method':'Same horizontal row indicates pada vedha; different rows within the same Nadi indicate cancellation. Original workbook scores remain unchanged.'}
 shutil.copyfile(nsrc,dest/'nadi-source.xlsx')
(dest/'scores.json').write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n')
shutil.copyfile(src,dest/'source.xlsx')
payload=json.dumps(obj,ensure_ascii=False,separators=(',',':')).replace('</','<\\/')
(ROOT/'vendor/matchmaking/workbook-data.js').write_text('/* Generated from the user-supplied workbook. */\nglobalThis.KPVedicWorkbookData='+payload+';\n')
print(json.dumps(obj['metadata'],ensure_ascii=False,indent=2))
