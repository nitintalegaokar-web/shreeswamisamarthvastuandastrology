"""Supplied workbook integrity, independent interpolation and native chart integration."""
import bisect
import hashlib
import json
import math
import re
import subprocess
import unittest
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import expect
import test_personal_predictions as personal
import test_protected_server as protected

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / 'vendor/houses/tables.json').read_text())
HOUSES = [10, 11, 12, 1, 2, 3]


def mix(a, b, fraction):
    return (a + ((b-a+180) % 360-180)*fraction) % 360


def reference(latitude, time, ayan, columns="numeric", round_ishta=True):
    pool=DATA["printedValues"] if columns=="printed" else DATA["values"]
    low = min(len(DATA['times'])-2, bisect.bisect_right(DATA['times'], time)-1)
    high = low+1
    f = (time-DATA['times'][low])/(DATA['times'][high]-DATA['times'][low])
    a, b = math.floor(latitude), math.ceil(latitude)
    if round_ishta:
        values=[]
        for i in range(6):
            small=mix(pool[a][low][i],pool[b][low][i],latitude-a)
            large=mix(pool[a][high][i],pool[b][high][i],latitude-a)
            arcseconds=((large-small+180)%360-180)*f*3600
            minutes=math.floor(abs(arcseconds)/60)+(1 if abs(arcseconds)%60>=50-1e-8 else 0)
            values.append((small+(-minutes if arcseconds<0 else minutes)/60)%360)
    else:
        values = [mix(mix(pool[a][low][i], pool[a][high][i], f), mix(pool[b][low][i], pool[b][high][i], f), latitude-a) for i in range(6)]
    cusps = [None]*12
    for h, x in zip(HOUSES, values):
        cusps[h-1] = (x-ayan) % 360
        cusps[(h+5) % 12] = (x-ayan+180) % 360
    return cusps


class TableTests(unittest.TestCase):
    def test_printed_house_degree_columns_and_ascendant_dm(self):
        script="const fs=require('fs'),lib=require('./vendor/houses/table.js'),data=JSON.parse(fs.readFileSync('./vendor/houses/tables.json')),t=lib.create(data);console.log(JSON.stringify([t.lookup(0,0,0,'printed'),t.lookup(0,220,0,'printed'),t.lookup(0,110,24,'printed')]));"
        result=json.loads(subprocess.check_output(['node','-e',script],cwd=ROOT))
        self.assertEqual(DATA['metadata']['printedSourceColumns'],['10H','11H','12H','Asc DM','2H','3H'])
        self.assertEqual([result[0]['tropicalCusps'][h-1] for h in HOUSES],[0,32,62,90,118,148])
        self.assertEqual([result[1]['tropicalCusps'][h-1] for h in [10,11,12,2,3]],[1,33,63,119,149])
        self.assertAlmostEqual(result[1]['tropicalCusps'][0],90+50/60,places=10)
        self.assertAlmostEqual(result[2]['cusps'][0],90+25/60-24,places=10)
        self.assertAlmostEqual(result[2]['cusps'][9],336.5,places=10)
        self.assertEqual(result[2]['tableColumns'],'printed')

    def test_every_uploaded_printed_workbook_row_is_preserved_and_hashed(self):
        rows, duplicates = 0, 0
        signs = ['Aries','Taurus','Gemini','Cancer','Leo','Virgo','Libra','Scorpio','Sagittarius','Capricorn','Aquarius','Pisces']
        self.assertFalse(DATA['metadata']['numericColumnsProvided'])
        self.assertEqual(DATA['metadata']['sourcePrecisionArcSeconds'],[3600,3600,3600,60,3600,3600])
        for source in DATA['metadata']['sources']:
            path = ROOT / 'vendor/houses/source' / source['filename']
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), source['sha256'])
            workbook = load_workbook(path, read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames, source['sheets'])
            for sheet in workbook:
                lat = int(sheet.title[:-1]); last = None; packed = []
                self.assertEqual(list(next(iter(sheet.values))),['ST HH:MM:SS','10H','11H','12H','Asc DM','2H','3H'])
                for row in list(sheet.values)[1:]:
                    h,m,s=map(int,row[0].split(':'));time=h*3600+m*60+s
                    values=[]
                    for column in row[1:]:
                        sign, number = column.split(' ')
                        parts=list(map(int,re.findall(r'\d+',number)))
                        values.append((signs.index(sign)*30+parts[0]+(parts[1]/60 if len(parts)>1 else 0))%360)
                    if time == last:
                        self.assertEqual(values, packed[-1][1]); duplicates += 1
                    else:
                        packed.append((time, values)); rows += 1
                    last = time
                self.assertEqual([r[0] for r in packed], DATA['times'])
                self.assertEqual([r[1] for r in packed], DATA['values'][lat])
                self.assertEqual([r[1] for r in packed], DATA['printedValues'][lat])
            workbook.close()
        self.assertEqual(rows, 22021); self.assertEqual(duplicates, 671)

    def test_exact_rows_all_latitudes_interpolation_wrap_and_range_checks(self):
        cases = [[0, 0, 0], [60, 86400, 24], [18.8, 43201, 23.85], [59.99, 86399.9, 24.001], [0.01, 219, 0], [37.25, 21598.5, 23.7]]
        js = """const fs=require('fs');const data=JSON.parse(fs.readFileSync('vendor/houses/tables.json'));const api=require('./vendor/houses/table.js').create(data);let maximum=0;for(let lat=0;lat<=60;lat++)for(let row=0;row<data.times.length;row++){const result=api.lookup(lat,data.times[row]);for(let i=0;i<6;i++)maximum=Math.max(maximum,Math.abs(((result.tropicalCusps[[10,11,12,1,2,3][i]-1]-data.values[lat][row][i]+540)%360)-180));}const cases=JSON.parse(process.argv[1]);const invalid=[[-1,0],[60.01,0],[0,-1],[0,86401],[null,0],[0,null],[0,0,-1]].map(x=>{try{api.lookup(...x);return false;}catch(_){return true;}});console.log(JSON.stringify({maximum,results:cases.map(x=>api.lookup(...x)),invalid}));"""
        result = subprocess.run(['node', '-e', js, json.dumps(cases)], cwd=ROOT, text=True, capture_output=True, check=True)
        actual = json.loads(result.stdout)
        self.assertLess(actual['maximum'], 1e-10); self.assertTrue(all(actual['invalid']))
        for inputs, result in zip(cases, actual['results']):
            for a, b in zip(result['cusps'], reference(*inputs,round_ishta=False)):
                self.assertAlmostEqual(a, b, places=10)
            for i in range(6):
                self.assertAlmostEqual((result['cusps'][i+6]-result['cusps'][i]) % 360, 180)
        self.assertAlmostEqual(actual['results'][0]['tropicalCusps'][9], 0)
        self.assertLess(abs((actual['results'][3]['tropicalCusps'][9]+180) % 360-180), .01)


class BhavsadhanTests(unittest.TestCase):
    setUpClass = classmethod(personal.PersonalPredictionTests.setUpClass.__func__)
    tearDownClass = classmethod(personal.PersonalPredictionTests.tearDownClass.__func__)
    setUp = personal.PersonalPredictionTests.setUp
    tearDown = personal.PersonalPredictionTests.tearDown

    def sample(self):
        self.page.wait_for_selector('#r5-house-source', state='attached')
        return self.page.evaluate('({data:KPBhavSadhan.getData(),cusps:currentKPModel.houses.map(h=>h.longitude),ready:currentKPModel.ready,source:document.getElementById("r5-house-source").value,st:document.getElementById("birthPlaceSiderealTime").value})')

    def verify(self, result):
        self.assertTrue(result['ready']); d = result['data']; self.assertEqual(d['source'], 'uploaded')
        expected = reference(d.get('tableLatitude', d['latitude']), d['siderealSeconds'], d['ayanamsha'], columns=d.get('tableColumns','numeric'))
        self.assertEqual(result['cusps'], [math.floor(x*3600+.5) % 1296000 for x in expected])
        self.assertEqual(result['source'], 'automatic')

    def test_latitude_minute_selection_changes_actual_table_calculations(self):
        p=self.page;self.sample()
        inputs=[0,0.5,0.5+1/3600,18+29/60+59/3600,18.5,18.5+1/3600,18+59/60+59/3600,19,59.5,59.5+1/3600,60,60.5,60.5+1/3600]
        self.assertEqual(p.evaluate('values=>values.map(KPBhavSadhan.selectLatitude)',inputs),[0,1,1,18,19,19,19,19,60,60,60,61,61])
        common={'instant-prashna':'0','dob':'1986-07-15','birthTime':'13:11:00','lon':'74:13:00','birth-utc-offset':'5.5','r5-house-source':'automatic'}
        planets=None
        for latitude,chosen in [('18:29:59',18),('18:30:00',19),('18:30:01',19),('18:59:59',19),('60:29:59',60)]:
            p.evaluate('fields=>{for(const [id,v] of Object.entries(fields))document.getElementById(id).value=v;calculateAll()}',{**common,'lat':latitude})
            sample=self.sample();self.verify(sample);data=sample['data'];b=data['brackets']
            actual=sum(int(x)/divisor for x,divisor in zip(latitude.split(':'),[1,60,3600]))
            self.assertEqual(p.locator('#lat').input_value(),latitude)
            self.assertAlmostEqual(data['latitude'],actual,places=12)
            self.assertEqual(data['tableLatitude'],chosen)
            self.assertEqual(data['latitudeRule'],'minutes-half-up')
            self.assertEqual((b['latitudeLow'],b['latitudeHigh'],b['latitudeFraction']),(chosen,chosen,0))
            row=DATA['times'].index(b['timeLow'])
            self.assertEqual(b['corners'][0],DATA['printedValues'][chosen][row])
            for actual_value,source_value in zip(b['small'],DATA['printedValues'][chosen][row]):self.assertAlmostEqual(actual_value,source_value,places=11)
            self.assertEqual(p.locator('#r5_small_12').input_value(),str(round(DATA['printedValues'][chosen][row][2])))
            expect(p.locator('#r5-house-status')).to_contain_text(f'Raphael table latitude {chosen}° N')
            expected=[math.floor(x*3600+.5)%1296000 for x in reference(chosen,data['siderealSeconds'],data['ayanamsha'],'printed')]
            self.assertEqual(sample['cusps'],expected)
            previous=[math.floor(x*3600+.5)%1296000 for x in reference(actual,data['siderealSeconds'],data['ayanamsha'],'printed')] if actual<=60 else None
            if previous is not None:self.assertNotEqual(sample['cusps'],previous)
            now=p.evaluate('currentKPModel.planets.map(x=>x.longitude)')
            if planets is None:planets=now
            else:self.assertEqual(now,planets)
        # No 61-degree sheet is invented or clamped to 60; fallback uses true latitude.
        p.evaluate('document.getElementById("lat").value="60:30:00";calculateAll()')
        data=self.sample()['data'];self.assertEqual(data['source'],'computed');self.assertAlmostEqual(data['latitude'],60.5)
        # Explicit astronomical mode keeps the true geographical latitude.
        p.locator("#quick-admin-calculations").click();p.locator("[name=admin-pin]").fill("684219");p.locator("[name=admin-confirm]").fill("684219");p.locator("#admin-pin-submit").click();expect(p.locator("#admin-calculation-host")).to_be_visible()
        p.evaluate('document.getElementById("r5-house-source").value="computed";document.getElementById("lat").value="18:30:01";calculateAll()')
        data=self.sample()['data'];self.assertEqual(data['source'],'computed');self.assertAlmostEqual(data['latitude'],18.5+1/3600)
        self.assertNotIn('tableLatitude',data)
        p.evaluate('document.getElementById("r5-house-source").value="automatic";calculateAll();KPPreferences.save({...KPPreferences.get(),language:"marathi"})')
        expect(p.locator('#r5-source-latitude-note')).to_contain_text('Excel साठी अक्षांश: 19°')
        expect(p.locator('#r5-source-latitude-note')).to_contain_text('३०′ किंवा अधिक')

    def test_source_rows_keep_excel_degrees_and_ascendant_dm(self):
        p=self.page
        p.wait_for_selector('#r5-source-latitude-details',state='attached')
        p.locator('#quick-admin-calculations').click()
        p.locator('[name="admin-pin"]').fill('684219');p.locator('[name="admin-confirm"]').fill('684219');p.locator('#admin-pin-submit').click()
        p.locator('[data-admin-worksheet="raphael5"]').click()
        def dm(x):
            minutes=math.floor(x*60+.5)%21600
            return f'{minutes//60}:{minutes%60:02}'
        def dms(x):
            n=math.floor(x*3600+.5)
            return f'{n//3600}:{n%3600//60:02}:{n%60:02}'
        for lat in ['18:47:00','18:48:00','0:00:00','60:00:00']:
            p.evaluate('lat=>{document.getElementById("lat").value=lat;document.getElementById("instant-prashna").value="0";document.getElementById("dob").value="1986-07-15";document.getElementById("birthTime").value="13:11:00";document.getElementById("lon").value="74:13:00";calculateAll()}',lat)
            result=self.sample();self.verify(result);b=result['data']['brackets']
            self.assertEqual(b['corners'][0],DATA['printedValues'][b['latitudeLow']][DATA['times'].index(b['timeLow'])])
            self.assertEqual(b['corners'][1],DATA['printedValues'][b['latitudeLow']][DATA['times'].index(b['timeHigh'])])
            for field,corner in [('small',0),('large',1)]:
                for i,h in enumerate(HOUSES):
                    text=p.locator(f'#r5_{field}_{h}').input_value()
                    self.assertEqual(text,dm(b['corners'][corner][i]) if h==1 else str(round(b['corners'][corner][i])))
                    self.assertRegex(text,r'^\d{1,3}:[0-5]\d$' if h==1 else r'^\d{1,3}$')
            for i,h in enumerate(HOUSES):
                self.assertEqual(p.locator(f'#r5_small2_{h}').input_value(),dms(b['small'][i]))
            extra=p.locator('#r5-source-latitude-details')
            if b['latitudeHigh']!=b['latitudeLow']:
                expect(extra).to_be_visible();p.locator('#r5-source-latitude-details summary').click()
                for field,corner in [('small',2),('large',3)]:
                    for i,h in enumerate(HOUSES):
                        self.assertEqual(p.locator(f'#r5-source-high-{field}-{h}').inner_text(),dm(b['corners'][corner][i]) if h==1 else str(round(b['corners'][corner][i])))
                p.locator('#r5-source-latitude-details summary').click()
            else:expect(extra).to_be_hidden()
        p.evaluate('document.getElementById("lat").value="18:47:00";calculateAll();KPPreferences.save({...KPPreferences.get(),language:"marathi"})')
        expect(p.locator('#r5-source-latitude-note')).to_contain_text('Excel साठी अक्षांश')
        expect(p.locator('#r5_small2_10').locator('..').locator('..')).to_contain_text('छोटी सां. वेळ')
        p.locator('#raphael5 .raphael5-table').screenshot(path='/tmp/kp-raphel-whole-degree-source-rows.png')

    def test_traditional_ten_six_st_and_birth_instant_kp_ayanamsha(self):
        def seconds(text):
            sign=-1 if text.startswith('-') else 1
            return sign*sum(int(value)*unit for value,unit in zip(text.lstrip('-+ ').split(':'),(3600,60,1)))
        def ten_six(elapsed):
            sign=-1 if elapsed<0 else 1;absolute=abs(elapsed)
            hours=absolute//3600;minutes=absolute%3600//60
            rounded=hours*10+minutes//6+(1 if minutes%6>=3 else 0)
            return sign*rounded
        cases=[('1986-07-15','13:11:00','18:48:00','74:13:00','5.5','0','82:30:00'),
               ('1986-07-15','01:00:00','18:48:00','74:13:00','5.5','0','82:30:00'),
               ('1987-01-01','01:00:00','18:48:00','74:13:00','5.5','0','82:30:00'),
               ('2024-02-29','23:59:59','18:48:00','74:13:00','5.5','0','82:30:00'),
               ('1986-07-15','23:59:59','40:42:00','-74:00:00','-5','60','-75:00:00')]
        for case in cases:
            values=dict(zip(['dob','birthTime','lat','lon','birth-utc-offset','birth-dst-minutes','stdLon'],case))
            result=self.page.evaluate('''values=>{for(const [id,value] of Object.entries({...values,'instant-prashna':'0'}))document.getElementById(id).value=value;calculateAll();const m=KPHorary.nativeMoment();return {date:m.date.toISOString(),house:KPBhavSadhan.getData(),cusps:currentKPModel.houses.map(h=>h.longitude),annual:KPAnnualAyanamsha.getTable().rows,st:document.getElementById('birthPlaceSiderealTime').value,worksheet:Object.fromEntries(['baseSidereal0530','stLmtDifference','stTenSixCorrection'].map(id=>[id,document.getElementById(id).value]))};}''',values)
            d=result['house'];worksheet=result['worksheet'];elapsed=seconds(worksheet['stLmtDifference']);correction=ten_six(elapsed)
            self.assertEqual(seconds(worksheet['stTenSixCorrection']),correction)
            expected=(seconds(worksheet['baseSidereal0530'])+elapsed+correction)%86400
            self.assertEqual(d['siderealSeconds'],expected,'Bhavsadhan must use the same traditional 10/6 ST shown by the worksheet.')
            instant=datetime.fromisoformat(result['date'].replace('Z','+00:00'));year=instant.year
            anchors={r['year']:r['dmsArcSeconds'] for r in result['annual']}
            start=datetime(year,1,1,tzinfo=timezone.utc);end=datetime(year+1,1,1,tzinfo=timezone.utc)
            ayan=(anchors[year]+(anchors[year+1]-anchors[year])*(instant-start)/(end-start))/3600
            self.assertAlmostEqual(d['ayanamsha'],ayan,places=11)
            expected_cusps=reference(d.get('tableLatitude', d['latitude']),d['siderealSeconds'],ayan,'printed')
            self.assertEqual(result['cusps'],[math.floor(x*3600+.5)%1296000 for x in expected_cusps])
            rounded=math.floor(d['siderealSeconds']+.5)
            self.assertEqual(result['st'],f'{rounded//3600}:{rounded%3600//60:02}:{rounded%60:02}')

    def test_automatic_native_instant_edits_fallback_and_invalid_location(self):
        p = self.page; self.verify(self.sample())
        p.evaluate('''()=>{for(const [id,v] of Object.entries({'instant-prashna':'0','dob':'1986-07-15','birthTime':'13:11:00','lat':'18:48:00','lon':'74:13:00','stdLon':'82:30:00','birth-utc-offset':'5.5','p6-ephemeris-source':'automatic','st-ephemeris-source':'automatic'}))document.getElementById(id).value=v;calculateAll();}''')
        p.wait_for_timeout(250); before = self.sample(); self.verify(before)
        d = before['data']; b = d['brackets']
        self.assertEqual(p.locator('#stSmallTime').input_value(), f"{b['timeLow']//3600}:{b['timeLow']%3600//60:02}:{b['timeLow']%60:02}")
        original_planets = p.evaluate('currentKPModel.planets.map(x=>x.longitude)')
        # An explicit source override is an administrator worksheet operation.
        p.locator('#quick-admin-calculations').click();p.locator('[name="admin-pin"]').fill('684219');p.locator('[name="admin-confirm"]').fill('684219');p.locator('#admin-pin-submit').click();expect(p.locator('#admin-calculation-host')).to_be_visible()
        p.evaluate('document.getElementById("r5-house-source").value="computed";calculateAll()'); self.assertEqual(self.sample()['data']['source'], 'computed'); self.assertEqual(p.evaluate('currentKPModel.planets.map(x=>x.longitude)'), original_planets)
        p.evaluate('document.getElementById("r5-house-source").value="automatic";document.getElementById("birthTime").value="04:03:02";document.getElementById("birthTime").dispatchEvent(new Event("change",{bubbles:true}))')
        p.wait_for_timeout(500); changed = self.sample(); self.verify(changed); self.assertNotEqual(changed['cusps'], before['cusps'])
        for latitude, source in [('60:00:00','uploaded'), ('60:30:00','computed'), ('60:30:01','computed'), ('-18:48:00','computed')]:
            p.evaluate('lat=>{document.getElementById("lat").value=lat;calculateAll()}', latitude); self.assertEqual(self.sample()['data']['source'], source); self.assertTrue(self.sample()['ready'])
        p.evaluate('document.getElementById("lat").value="70:00:00";calculateAll()'); self.assertFalse(self.sample()['ready']); self.assertIsNone(self.sample()['data']); self.assertEqual(p.locator('#r5_nirayan_1').input_value(), '')
        self.assertEqual(p.evaluate('Array.from({length:3},()=>KPVedic.core().ready)'), [False]*3)
        p.evaluate('document.getElementById("lat").value="18:48:00";calculateAll()'); self.verify(self.sample()); self.assertTrue(p.evaluate('KPVedic.core().ready'))

    def test_explicit_manual_preservation_horary_and_administrator_source(self):
        p = self.page; self.sample(); p.locator('#quick-admin-calculations').click(); p.locator('[name="admin-pin"]').fill('684219'); p.locator('[name="admin-confirm"]').fill('684219'); p.locator('#admin-pin-submit').click(); expect(p.locator('#admin-calculation-host')).to_be_visible(); p.locator('[data-admin-worksheet="raphael5"]').click()
        expect(p.locator('#r5-house-source')).to_be_visible(); p.screenshot(path='/tmp/uploaded-bhav-tables.png')
        p.evaluate('document.getElementById("instant-prashna").value="0";calculateAll()')
        p.locator('#r5_small_1').fill('100:12:34'); self.assertEqual(p.locator('#r5-house-source').input_value(), 'manual'); self.assertEqual(p.locator('#r5_small_1').input_value(), '100:12:34')
        p.evaluate('calculateAll()'); self.assertEqual(p.locator('#r5_small_1').input_value(), '100:12:34')
        saved = p.evaluate('getChartData()'); self.assertEqual(saved['fields']['r5-house-source']['value'], 'manual')
        legacy = json.loads(json.dumps(saved)); legacy.pop('manualCalculations', None); del legacy['fields']['r5-house-source']
        p.evaluate('chart=>{document.getElementById("r5-house-source").value="automatic";restoreChartData(chart)}', legacy)
        self.assertEqual(p.locator('#r5-house-source').input_value(), 'automatic'); self.assertNotEqual(p.locator('#r5_small_1').input_value(), '100:12:34')
        # Only an explicit ADM Save-manual record may restore manual values.
        p.locator('#r5_small_1').fill('100:12:34')
        p.evaluate('''()=>{const record=KPAutomaticCalculations.capture();KPAutomaticCalculations.setSaved(record);KPAutomaticCalculations.persist();}''')
        explicit = p.evaluate('getChartData()'); self.assertIn('manualCalculations', explicit)
        p.evaluate('chart=>{document.getElementById("r5-house-source").value="automatic";restoreChartData(chart)}', explicit)
        self.assertEqual(p.locator('#r5-house-source').input_value(), 'manual'); self.assertEqual(p.locator('#r5_small_1').input_value(), '100:12:34')
        p.evaluate('KPAutomaticCalculations.setSaved(null);KPAutomaticCalculations.persist()')
        p.evaluate('document.getElementById("r5-house-source").value="automatic";document.getElementById("chart-kind").value="horary";document.getElementById("horary-number").value="123";calculateAll()')
        p.wait_for_timeout(250); result = p.evaluate('({native:currentKPModel.houses.map(h=>h.longitude),horary:KPHorary.getData().cusps.map(c=>c*3600),source:KPBhavSadhan.getData()})'); self.assertIsNone(result['source'])
        for a, b in zip(result['native'], result['horary']): self.assertAlmostEqual(a, b)
        p.evaluate('KPPreferences.save({...KPPreferences.get(),language:"marathi"})'); expect(p.locator('#r5-house-source option').first).to_have_text('अपलोड केलेल्या सारण्या · स्वयंचलित')

    def test_offline_embedded_tables_and_saved_automatic_source(self):
        p = self.page
        # The managed Chromium policy blocks file://. Start with a clean HTTP
        # origin, then disconnect networking before loading the standalone HTML.
        p.goto(self.url.rsplit('/',1)[0]+'/offline-check')
        self.context.set_offline(True)
        p.set_content((ROOT / 'index.html').read_text(), wait_until='load', timeout=60000)
        self.verify(self.sample())
        saved = p.evaluate('getChartData()'); self.assertEqual(saved['fields']['r5-house-source']['value'], 'automatic')
        original = self.sample()['cusps']
        p.evaluate('chart=>{document.getElementById("r5-house-source").value="computed";restoreChartData(chart)}', saved)
        self.verify(self.sample()); self.assertEqual(self.sample()['cusps'], original)


class HousePrivateTests(unittest.TestCase):
    setUpClass = classmethod(protected.ProtectedServerTests.setUpClass.__func__)
    tearDownClass = classmethod(protected.ProtectedServerTests.tearDownClass.__func__)
    setUp = protected.ProtectedServerTests.setUp
    tearDown = protected.ProtectedServerTests.tearDown
    private = protected.ProtectedServerTests.private
    post = protected.ProtectedServerTests.post

    def test_private_worker_embeds_tables_and_hides_house_controls(self):
        result = self.private('({data:KPBhavSadhan.getData(),cusps:currentKPModel.houses.map(h=>h.longitude),ready:currentKPModel.ready})')
        self.assertTrue(result['ready']); d = result['data']; self.assertEqual(d['source'], 'uploaded')
        self.assertEqual(result['cusps'], [math.floor(x*3600+.5) % 1296000 for x in reference(d.get('tableLatitude', d['latitude']), d['siderealSeconds'], d['ayanamsha'], columns=d.get('tableColumns','numeric'))])
        self.assertEqual(self.page.evaluate('typeof KPBhavSadhan'), 'undefined')
        self.assertEqual(self.page.locator('#r5-house-source,#kp-uploaded-house-data').count(), 0)
        expect(self.page.locator('#quick-admin-calculations')).to_be_hidden()
        chart = self.page.request.get(self.url+'/export').json()
        self.assertNotIn('r5-house-source', chart['fields'])
        chart['fields']['instant-prashna']['value'] = '0'
        chart['fields']['dob']['value'] = '1986-07-15'
        chart['fields']['birthTime']['value'] = '13:11:00'
        chart['fields']['lat']['value'] = '18:48:00'
        chart['fields']['lon']['value'] = '74:13:00'
        # Private administrator overrides are not accepted through hosted imports.
        chart['manualCalculations']={'version':'not-a-private-command'}
        response=self.post('/import', {'chart':chart});self.assertEqual(response.status, 200,response.text())
        imported = self.private('({data:KPBhavSadhan.getData(),cusps:currentKPModel.houses.map(h=>h.longitude),ready:currentKPModel.ready,source:document.getElementById("r5-house-source").value})')
        self.assertTrue(imported['ready']); self.assertEqual(imported['source'], 'automatic')
        d = imported['data']; self.assertEqual(d['source'], 'uploaded')
        self.assertEqual(imported['cusps'], [math.floor(x*3600+.5) % 1296000 for x in reference(d.get('tableLatitude', d['latitude']), d['siderealSeconds'], d['ayanamsha'], columns=d.get('tableColumns','numeric'))])
